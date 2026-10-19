"""Write results into the sheet's own Result and Evidence / Page Ref columns, plus extras."""
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

FILL = {"PASS": "C6EFCE", "FAIL": "FFC7CE", "UNCLEAR": "FFEB9C"}


def fmt_evidence(r):
    out = []
    for c in r.citations:
        pg = f" p.{c['page']}" if c.get("page") else ""
        out.append(f"{c.get('document', '')}{pg}: \"{c.get('quote', '')}\"".strip())
    return "\n".join(out)


def write_results(sheet_df, tests, results, path, model_name, as_of, cfg):
    df = sheet_df.copy()
    by_id = {r.test_id: r for r in results}
    cols = ["Result", "Evidence / Page Ref", "Reason", "Review Flag", "Documents Used", "Mode", "Notes"]
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    id_col = next(c for c in df.columns if str(c).strip().lower() in ("test name", "test id", "test"))
    for i, row in df.iterrows():
        r = by_id.get(str(row[id_col]).strip())
        if not r:
            continue
        df.at[i, "Result"] = r.verdict
        df.at[i, "Evidence / Page Ref"] = fmt_evidence(r) + (f"\nAnswer: {r.answer}" if r.answer else "")
        df.at[i, "Reason"] = r.reason
        df.at[i, "Review Flag"] = "YES" if r.review_flag else ""
        df.at[i, "Documents Used"] = ", ".join(r.documents)
        df.at[i, "Mode"] = r.mode
        df.at[i, "Notes"] = "; ".join(r.notes)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, index=False, sheet_name="Results")
        counts = pd.Series([r.verdict for r in results]).value_counts()
        info = pd.DataFrame({"Item": ["Run time", "Model", "As-of date", "Tests run", "PASS", "FAIL", "UNCLEAR",
                                      "Full-document limit (chars)", "Top-k per document"],
                             "Value": [datetime.now().isoformat(timespec="seconds"), model_name, as_of.isoformat(),
                                       len(results), int(counts.get("PASS", 0)), int(counts.get("FAIL", 0)),
                                       int(counts.get("UNCLEAR", 0)), cfg["retrieval"]["full_doc_char_limit"],
                                       cfg["retrieval"]["top_k"]]})
        info.to_excel(xw, index=False, sheet_name="Run info")
    wb = load_workbook(path)
    ws = wb["Results"]
    hdr = {c.value: c.column for c in ws[1]}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0D2748")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        v = row[hdr["Result"] - 1].value
        if v in FILL:
            row[hdr["Result"] - 1].fill = PatternFill("solid", fgColor=FILL[v])
    for name, w in {"Evidence / Page Ref": 55, "Reason": 50, "Notes": 40, "Tests": 60, "File Name": 34}.items():
        if name in hdr:
            ws.column_dimensions[ws.cell(1, hdr[name]).column_letter].width = w
    ws.freeze_panes = "B2"
    wb.save(path)
